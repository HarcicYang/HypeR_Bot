"""Local document extraction for QQ uploads and downloadable files."""

from __future__ import annotations

import csv
import dataclasses
import io
import json
import os
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, cast

from hyperot import configurator
from typing_extensions import override

config = configurator.BotConfig.get("hyper-bot")


def _config_int(name: str, default: int) -> int:
    value = config.others.get(name)
    if value is None:
        return default
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


MAX_DOCUMENT_BYTES = max(1, _config_int("agent_document_max_mb", 50)) * 1024 * 1024
MAX_ARCHIVE_BYTES = max(1, _config_int("agent_archive_max_mb", 100)) * 1024 * 1024
MAX_TEXT_CHARS = max(0, _config_int("agent_document_max_chars", 10_000_000))
VISION_PAGES = max(0, _config_int("agent_document_vision_pages", 5))
VISION_TEXT_THRESHOLD = max(1, _config_int("agent_document_vision_threshold", 200))

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".log",
    ".py",
    ".pyi",
    ".js",
    ".mjs",
    ".ts",
    ".tsx",
    ".jsx",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".xml",
    ".sql",
    ".sh",
    ".zsh",
    ".fish",
    ".bat",
    ".ps1",
    ".c",
    ".h",
    ".cc",
    ".cpp",
    ".hpp",
    ".java",
    ".kt",
    ".rs",
    ".go",
    ".rb",
    ".php",
    ".css",
    ".scss",
    ".less",
}
CSV_EXTENSIONS = {".csv", ".tsv"}
HTML_EXTENSIONS = {".html", ".htm", ".xhtml"}
ARCHIVE_EXTENSIONS = {".zip"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
LEGACY_OFFICE_EXTENSIONS = {".doc", ".xls", ".ppt"}
DOCX_EXTENSIONS = {".docx", ".docm"}
XLSX_EXTENSIONS = {".xlsx", ".xlsm"}
PPTX_EXTENSIONS = {".pptx", ".pptm"}


@dataclasses.dataclass(frozen=True)
class ParsedDocument:
    text: str
    kind: str
    metadata: dict[str, Any] = dataclasses.field(default_factory=dict)


class DocumentReadError(RuntimeError):
    """The document could not be downloaded or decoded."""


def _clip(text: str, limit: int = MAX_TEXT_CHARS) -> tuple[str, bool]:
    if limit <= 0 or len(text) <= limit:
        return text, False
    return text[:limit], True


def _decode_text(data: bytes) -> str:
    from charset_normalizer import from_bytes

    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    match = from_bytes(data).best()
    if match is None:
        return data.decode("utf-8", errors="replace")
    return str(match)


def _format_json(text: str) -> str:
    try:
        return json.dumps(json.loads(text), ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return text


def _format_csv(text: str, delimiter: str | None = None) -> str:
    lines = text.splitlines()
    if not lines:
        return text
    if delimiter is None:
        delimiter = "\t" if "\t" in lines[0] else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _check_zip_expansion(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        return
    if total > MAX_ARCHIVE_BYTES:
        raise DocumentReadError("文档解压后大小超过读取上限")


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    @override
    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"p", "div", "section", "article", "header", "footer", "li", "tr", "br", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    @override
    def handle_endtag(self, tag: str) -> None:
        if tag in {"p", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    @override
    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)
            self.parts.append(" ")

    def text(self) -> str:
        return "".join(self.parts)


def _parse_pdf(data: bytes) -> ParsedDocument:
    import pymupdf

    document = pymupdf.open(stream=data, filetype="pdf")
    try:
        parts: list[str] = []
        for index, page in enumerate(document, start=1):
            text = str(page.get_text("text", sort=True) or "").strip()
            if text:
                parts.append(f"## 第 {index} 页\n{text}")
        merged, truncated = _clip("\n\n".join(parts))
        return ParsedDocument(
            merged,
            "PDF",
            {
                "pages": document.page_count,
                "text_pages": len(parts),
                "scanned": len(merged) < VISION_TEXT_THRESHOLD,
                "truncated": truncated,
            },
        )
    finally:
        document.close()


def _parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document

    _check_zip_expansion(data)
    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = str(paragraph.style.name if paragraph.style is not None else "")
        if style.startswith("Heading "):
            level = style.removeprefix("Heading ").strip()
            prefix = "#" * min(int(level), 6) if level.isdigit() else "##"
            parts.append(f"{prefix} {text}")
        else:
            parts.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells) for row in table.rows]
        if rows:
            parts.append(f"## 表格 {table_index}\n" + "\n".join(rows))
    merged, truncated = _clip("\n\n".join(parts))
    return ParsedDocument(
        merged,
        "DOCX",
        {
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "truncated": truncated,
        },
    )


def _parse_xlsx(data: bytes) -> ParsedDocument:
    import openpyxl

    _check_zip_expansion(data)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        parts: list[str] = []
        row_count = 0
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    rows.append(" | ".join(cells))
                    row_count += 1
            if rows:
                parts.append(f"## 工作表: {sheet.title}\n" + "\n".join(rows))
        merged, truncated = _clip("\n\n".join(parts))
        return ParsedDocument(
            merged,
            "XLSX",
            {
                "sheets": workbook.sheetnames,
                "rows": row_count,
                "truncated": truncated,
            },
        )
    finally:
        workbook.close()


def _parse_pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation

    _check_zip_expansion(data)
    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            text = ""
            shape_any = cast(Any, shape)
            if getattr(shape_any, "has_text_frame", False):
                text = str(shape_any.text or "").strip()
            if text:
                slide_parts.append(text)
            if getattr(shape_any, "has_table", False):
                rows = [
                    " | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells)
                    for row in shape_any.table.rows
                ]
                if rows:
                    slide_parts.append("\n".join(rows))
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            notes = (notes_frame.text if notes_frame is not None else "").strip()
            if notes:
                slide_parts.append(f"备注:\n{notes}")
        if slide_parts:
            parts.append(f"## 幻灯片 {slide_index}\n" + "\n\n".join(slide_parts))
    merged, truncated = _clip("\n\n".join(parts))
    return ParsedDocument(
        merged,
        "PPTX",
        {"slides": len(presentation.slides), "truncated": truncated},
    )


def _parse_zip(data: bytes, member: str = "") -> ParsedDocument:
    archive = zipfile.ZipFile(io.BytesIO(data))
    try:
        if member:
            try:
                info = archive.getinfo(member)
            except KeyError as exc:
                raise DocumentReadError(f"压缩包中不存在文件: {member}") from exc
            if info.file_size > MAX_DOCUMENT_BYTES:
                raise DocumentReadError("压缩包内文件超过读取上限")
            content = archive.read(info)
            return parse_document(content, info.filename)

        total = sum(info.file_size for info in archive.infolist())
        if total > MAX_ARCHIVE_BYTES:
            raise DocumentReadError("压缩包解压后大小超过读取上限")
        lines = ["压缩包内容:"]
        for info in archive.infolist():
            kind = "目录" if info.is_dir() else "文件"
            lines.append(f"- [{kind}] {info.filename} ({info.file_size} bytes)")
        text, truncated = _clip("\n".join(lines))
        return ParsedDocument(
            text,
            "ZIP",
            {"members": len(archive.infolist()), "uncompressed_bytes": total, "truncated": truncated},
        )
    finally:
        archive.close()


def render_pdf_pages(data: bytes, limit: int = VISION_PAGES) -> list[bytes]:
    """Render leading PDF pages as PNG for the optional vision fallback."""
    if limit <= 0:
        return []
    import pymupdf

    document = pymupdf.open(stream=data, filetype="pdf")
    try:
        images: list[bytes] = []
        for index in range(min(limit, document.page_count)):
            page = document.load_page(index)
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
            images.append(pixmap.tobytes("png"))
        return images
    finally:
        document.close()


def parse_document(data: bytes, filename: str, member: str = "") -> ParsedDocument:
    """Extract readable text from bytes based on file type."""
    if len(data) > MAX_DOCUMENT_BYTES:
        raise DocumentReadError(f"文件超过读取上限 ({MAX_DOCUMENT_BYTES // 1024 // 1024} MB)")
    name = os.path.basename(filename or "document")
    extension = Path(name).suffix.lower()

    if extension == ".pdf" or data.startswith(b"%PDF"):
        return _parse_pdf(data)
    if extension in DOCX_EXTENSIONS:
        return _parse_docx(data)
    if extension in XLSX_EXTENSIONS:
        return _parse_xlsx(data)
    if extension in PPTX_EXTENSIONS:
        return _parse_pptx(data)
    if extension in ARCHIVE_EXTENSIONS:
        return _parse_zip(data, member)
    if extension in LEGACY_OFFICE_EXTENSIONS:
        raise DocumentReadError(f"暂不支持旧版 Office 格式: {extension}")
    if extension in IMAGE_EXTENSIONS:
        return ParsedDocument("", "IMAGE", {"vision_required": True})

    text = _decode_text(data)
    if extension in HTML_EXTENSIONS:
        parser = _TextHTMLParser()
        parser.feed(text)
        text = parser.text()
        kind = "HTML"
    elif extension == ".json":
        text = _format_json(text)
        kind = "JSON"
    elif extension in CSV_EXTENSIONS:
        text = _format_csv(text, "\t" if extension == ".tsv" else None)
        kind = "CSV"
    elif extension in TEXT_EXTENSIONS:
        kind = "TEXT"
    else:
        printable = sum(character.isprintable() or character in "\n\r\t" for character in text)
        if text and printable / len(text) < 0.85:
            raise DocumentReadError(f"暂不支持的文件格式: {extension or 'unknown'}")
        kind = "TEXT"

    clipped, truncated = _clip(text)
    return ParsedDocument(clipped, kind, {"truncated": truncated})
